"""Extract plain text from uploaded resume / job-description files."""
from __future__ import annotations

import io

SUPPORTED_EXTENSIONS = ("pdf", "docx", "txt", "md")


def extract_text(filename: str, data: bytes) -> str:
    """Return text from PDF, DOCX, TXT or MD bytes. Raises ValueError on bad input."""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type '.{ext}'. Use: {', '.join(SUPPORTED_EXTENSIONS)}.")
    if ext == "pdf":
        return _read_pdf(data)
    if ext == "docx":
        return _read_docx(data)
    return data.decode("utf-8", errors="ignore")


def _read_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            reader.decrypt("")
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception as exc:  # corrupted / unreadable
        raise ValueError(f"Could not read PDF: {exc}") from exc
    if not text.strip():
        raise ValueError("No selectable text found in the PDF (it may be a scanned image). Paste the text instead.")
    return text


def _read_docx(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ValueError(f"Could not read DOCX: {exc}") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)
